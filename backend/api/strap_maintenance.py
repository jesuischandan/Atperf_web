"""
STRAP Maintenance API — datasets, import, scoring cockpit, replay, backtest, expert decisions.

Access: active Atperf staff see every dataset; users of a company with the
`maintenance_pilot` flag see their own company's datasets only.
"""
import io
import json
import shutil
import uuid
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Literal, Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

import models
from database import get_gcp_db
from maintenance import engine as E
from maintenance import importer
from security import Principal, get_principal, rate_limit

router = APIRouter(prefix="/api/maintenance", tags=["STRAP Maintenance"])

UPLOAD_DIR = Path("uploads") / "maintenance"
SAMPLE_FILE = Path(__file__).resolve().parent.parent / "maintenance" / "sample" / "STRAP_Maintenance_Donnees_Fictives.xlsx"
MAX_UPLOAD_BYTES = 20 * 1024 * 1024


# ── Access ────────────────────────────────────────────────────────────────────
@dataclass
class Access:
    principal: Principal
    is_staff: bool
    company_id: Optional[str]

    @property
    def actor(self) -> tuple[str, str, str]:
        if self.principal.staff is not None:
            s = self.principal.staff
            return "staff", s.staff_id, f"{s.first_name} {s.last_name}"
        u = self.principal.user
        return "user", u.user_id, f"{u.first_name} {u.last_name}"


def maintenance_access(principal: Principal = Depends(get_principal), gcp_db: Session = Depends(get_gcp_db)) -> Access:
    if principal.staff is not None:
        return Access(principal, True, None)
    if principal.user is None:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Not authenticated. Please log in.")
    company = gcp_db.query(models.GCPCompany).filter(models.GCPCompany.company_id == principal.user.company_id).first()
    if not company or not company.maintenance_pilot:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="STRAP Maintenance is not enabled for your company.")
    return Access(principal, False, company.company_id)


def get_dataset(dataset_id: str, access: Access, gcp_db: Session) -> models.MaintDataset:
    ds = gcp_db.query(models.MaintDataset).filter(models.MaintDataset.id == dataset_id).first()
    if ds is None or (not access.is_staff and ds.company_id != access.company_id):
        raise HTTPException(status_code=404, detail="Dataset not found.")
    return ds


# ── Engine loading ────────────────────────────────────────────────────────────
def load_engine(ds: models.MaintDataset, gcp_db: Session) -> E.Engine:
    eqs = gcp_db.query(models.MaintEquipment).filter_by(dataset_id=ds.id).all()
    ns = gcp_db.query(models.MaintNotice).filter_by(dataset_id=ds.id).all()
    wos = gcp_db.query(models.MaintWorkOrder).filter_by(dataset_id=ds.id).all()
    return E.Engine(
        [E.Equipment(e.code, e.name or "", e.tag or "", e.subsystem or "", e.criticality_label or "", e.criticality_note) for e in eqs],
        [E.Notice(n.code, n.date, n.equipment, n.flux, n.symptom or "", n.description or "", n.category or "", n.status or "",
                  n.postponements or 0, n.issuer or "", n.context_note, dict(n.overrides or {})) for n in ns],
        [E.WorkOrder(o.code, o.date, o.equipment, o.flux, o.type or "", o.symptom or "", o.cause or "", o.downtime_hours,
                     o.status or "", o.linked_notice) for o in wos],
        E.Config.from_dicts(ds.weights, ds.thresholds, ds.windows),
    )


def parse_as_of(as_of: Optional[str], eng: E.Engine) -> date:
    if not as_of:
        return eng.study_date()
    try:
        return date.fromisoformat(as_of)
    except ValueError:
        raise HTTPException(status_code=400, detail="as_of must be YYYY-MM-DD")


def equipment_json(eng: E.Engine, code: str) -> dict:
    e = eng.equipment(code)
    return {"code": e.code, "name": e.name, "tag": e.tag or e.code, "subsystem": e.subsystem,
            "criticality": e.criticality_label, "criticality_note": E.criticality_note(e.criticality_label, e.criticality_note)}


def evidence_json(ev: E.Evidence) -> dict:
    return {"note": ev.note, "reason": ev.reason, "detail": ev.detail, "overridden": ev.overridden}


def latest_decisions(ds_id: str, gcp_db: Session, target_type: str) -> dict[str, models.MaintDecision]:
    rows = gcp_db.query(models.MaintDecision).filter_by(dataset_id=ds_id, target_type=target_type) \
        .order_by(models.MaintDecision.created_at, models.MaintDecision.id).all()
    return {r.target: r for r in rows}


def decision_json(d: Optional[models.MaintDecision]) -> Optional[dict]:
    if d is None:
        return None
    return {"id": d.id, "target_type": d.target_type, "target": d.target, "decision": d.decision, "comment": d.comment,
            "changes": d.changes, "score_snapshot": d.score_snapshot, "actor_name": d.actor_name, "actor_kind": d.actor_kind,
            "created_at": d.created_at.isoformat() if d.created_at else None}


def log_decision(gcp_db: Session, ds: models.MaintDataset, access: Access, target_type: str, target: str, decision: str,
                 comment: Optional[str] = None, changes: Optional[dict] = None, score: Optional[int] = None) -> models.MaintDecision:
    kind, actor_id, name = access.actor
    row = models.MaintDecision(dataset_id=ds.id, target_type=target_type, target=target, decision=decision,
                               comment=(comment or "").strip() or None, changes=changes, score_snapshot=score,
                               actor_kind=kind, actor_id=actor_id, actor_name=name)
    gcp_db.add(row)
    return row


# ── Datasets ──────────────────────────────────────────────────────────────────
def dataset_json(ds: models.MaintDataset, company_names: dict, counts: Optional[dict] = None) -> dict:
    return {"id": ds.id, "name": ds.name, "company_id": ds.company_id, "company_name": company_names.get(ds.company_id),
            "source_filename": ds.source_filename, "weights": ds.weights, "thresholds": ds.thresholds, "windows": ds.windows,
            "import_report": ds.import_report, "created_by_name": ds.created_by_name,
            "created_at": ds.created_at.isoformat() if ds.created_at else None, "counts": counts}


@router.get("/datasets")
def list_datasets(access: Access = Depends(maintenance_access), gcp_db: Session = Depends(get_gcp_db)):
    q = gcp_db.query(models.MaintDataset)
    if not access.is_staff:
        q = q.filter(models.MaintDataset.company_id == access.company_id)
    names = {c.company_id: c.company_name for c in gcp_db.query(models.GCPCompany).all()}
    return [dataset_json(ds, names, (ds.import_report or {}).get("counts")) for ds in q.order_by(models.MaintDataset.created_at.desc()).all()]


@router.get("/companies")
def pilot_companies(access: Access = Depends(maintenance_access), gcp_db: Session = Depends(get_gcp_db)):
    """Companies a dataset can be attached to (staff: all pilot companies; client: their own)."""
    q = gcp_db.query(models.GCPCompany).filter(models.GCPCompany.maintenance_pilot == 1)
    if not access.is_staff:
        q = q.filter(models.GCPCompany.company_id == access.company_id)
    return [{"company_id": c.company_id, "company_name": c.company_name} for c in q.order_by(models.GCPCompany.company_name).all()]


@router.get("/datasets/{dataset_id}")
def read_dataset(dataset_id: str, access: Access = Depends(maintenance_access), gcp_db: Session = Depends(get_gcp_db)):
    ds = get_dataset(dataset_id, access, gcp_db)
    names = {c.company_id: c.company_name for c in gcp_db.query(models.GCPCompany).filter_by(company_id=ds.company_id).all()}
    eng = load_engine(ds, gcp_db)
    out = dataset_json(ds, names, (ds.import_report or {}).get("counts"))
    out["study_date"] = eng.study_date().isoformat()
    out["can_delete"] = access.is_staff or access.actor[1] == ds.created_by_id
    return out


@router.delete("/datasets/{dataset_id}", status_code=204)
def delete_dataset(dataset_id: str, access: Access = Depends(maintenance_access), gcp_db: Session = Depends(get_gcp_db)):
    ds = get_dataset(dataset_id, access, gcp_db)
    if not (access.is_staff or access.actor[1] == ds.created_by_id):
        raise HTTPException(status_code=403, detail="Only Atperf staff or the person who imported it can delete this dataset.")
    for model in (models.MaintDecision, models.MaintLink, models.MaintNotice, models.MaintWorkOrder, models.MaintEquipment):
        gcp_db.query(model).filter(model.dataset_id == ds.id).delete()
    gcp_db.delete(ds)
    gcp_db.commit()


class ConfigUpdate(BaseModel):
    weights: dict[str, float]
    thresholds: dict[str, float]
    windows: dict[str, float] = Field(default_factory=dict)
    comment: Optional[str] = None


@router.put("/datasets/{dataset_id}/config")
def update_config(dataset_id: str, payload: ConfigUpdate, access: Access = Depends(maintenance_access),
                  gcp_db: Session = Depends(get_gcp_db)):
    ds = get_dataset(dataset_id, access, gcp_db)
    windows = {**(ds.windows or E.DEFAULT_WINDOWS), **payload.windows}
    errors = E.validate_config(payload.weights, payload.thresholds, windows)
    if errors:
        raise HTTPException(status_code=400, detail="; ".join(errors))
    before = {"weights": ds.weights, "thresholds": ds.thresholds, "windows": ds.windows}
    ds.weights = {k: int(v) if float(v).is_integer() else v for k, v in payload.weights.items()}
    ds.thresholds = {k: int(v) for k, v in payload.thresholds.items()}
    ds.windows = {k: int(v) for k, v in windows.items()}
    log_decision(gcp_db, ds, access, "config", "config", "updated", payload.comment,
                 {"before": before, "after": {"weights": ds.weights, "thresholds": ds.thresholds, "windows": ds.windows}})
    gcp_db.commit()
    return {"weights": ds.weights, "thresholds": ds.thresholds, "windows": ds.windows}


# ── Import ────────────────────────────────────────────────────────────────────
def _token_dir(token: str) -> Path:
    if not token or any(c not in "0123456789abcdef" for c in token):
        raise HTTPException(status_code=400, detail="Invalid import token.")
    return UPLOAD_DIR / token


def _load_token(token: str, access: Access) -> tuple[Path, list[importer.Source]]:
    folder = _token_dir(token)
    meta_path = folder / "meta.json"
    if not meta_path.exists():
        raise HTTPException(status_code=404, detail="Import session expired. Please upload the file again.")
    meta = json.loads(meta_path.read_text())
    if meta.get("owner") != access.actor[1]:
        raise HTTPException(status_code=404, detail="Import session expired. Please upload the file again.")
    files = [(name, (folder / f"{i}").read_bytes()) for i, name in enumerate(meta["files"])]
    return folder, importer.read_sources(files)


@router.post("/import/preview", dependencies=[Depends(rate_limit("maint-import", 30, 3600))])
async def import_preview(files: list[UploadFile] = File(...), access: Access = Depends(maintenance_access)):
    if not files or len(files) > 4:
        raise HTTPException(status_code=400, detail="Upload one Excel workbook or up to four CSV files.")
    token = uuid.uuid4().hex
    folder = UPLOAD_DIR / token
    folder.mkdir(parents=True, exist_ok=True)
    names, payload = [], []
    try:
        for i, f in enumerate(files):
            content = await f.read()
            if len(content) > MAX_UPLOAD_BYTES:
                raise HTTPException(status_code=400, detail="File too large (20 MB max).")
            name = Path(f.filename or f"file{i}").name
            if not name.lower().endswith((".xlsx", ".xlsm", ".xls", ".csv", ".txt")):
                raise HTTPException(status_code=400, detail="Use .xlsx or .csv files.")
            (folder / f"{i}").write_bytes(content)
            names.append(name)
            payload.append((name, content))
        (folder / "meta.json").write_text(json.dumps({"owner": access.actor[1], "files": names}))
        sources = importer.read_sources(payload)
        if not sources:
            raise HTTPException(status_code=400, detail="The file contains no data.")
        return {"token": token, "files": names, **importer.preview(sources)}
    except HTTPException:
        shutil.rmtree(folder, ignore_errors=True)
        raise
    except Exception as exc:  # unreadable workbook / CSV
        shutil.rmtree(folder, ignore_errors=True)
        raise HTTPException(status_code=400, detail=f"Could not read the file: {exc}")


class TableSpec(BaseModel):
    source: Optional[int] = None
    mapping: dict[str, Optional[str]] = Field(default_factory=dict)


class ImportCommit(BaseModel):
    token: str
    name: str = Field(min_length=1, max_length=255)
    company_id: Optional[str] = None
    tables: dict[Literal["equipments", "notices", "work_orders", "links"], TableSpec]


def _create_dataset(gcp_db: Session, access: Access, name: str, company_id: Optional[str], filename: str,
                    result: importer.ImportResult) -> models.MaintDataset:
    if access.is_staff:
        if company_id:
            company = gcp_db.query(models.GCPCompany).filter_by(company_id=company_id).first()
            if not company:
                raise HTTPException(status_code=400, detail="Unknown company.")
    else:
        company_id = access.company_id
    kind, actor_id, actor_name = access.actor
    ds = models.MaintDataset(name=name.strip(), company_id=company_id or None, source_filename=filename,
                             weights=dict(E.DEFAULT_WEIGHTS), thresholds=dict(E.DEFAULT_THRESHOLDS), windows=dict(E.DEFAULT_WINDOWS),
                             import_report=result.report(), created_by_kind=kind, created_by_id=actor_id, created_by_name=actor_name)
    gcp_db.add(ds)
    gcp_db.flush()
    gcp_db.add_all([models.MaintEquipment(dataset_id=ds.id, **e) for e in result.equipments])
    gcp_db.add_all([models.MaintNotice(dataset_id=ds.id, **n) for n in result.notices])
    gcp_db.add_all([models.MaintWorkOrder(dataset_id=ds.id, **o) for o in result.work_orders])
    gcp_db.add_all([models.MaintLink(dataset_id=ds.id, notice_code=l["notice"], wo_code=l["work_order"],
                                     expert_status=l["expert_status"] or None) for l in result.links])
    gcp_db.commit()
    return ds


@router.post("/import/commit", status_code=201)
def import_commit(payload: ImportCommit, access: Access = Depends(maintenance_access), gcp_db: Session = Depends(get_gcp_db)):
    folder, sources = _load_token(payload.token, access)
    try:
        result = importer.build(sources, {t: spec.model_dump() for t, spec in payload.tables.items()})
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    filename = ", ".join(json.loads((folder / "meta.json").read_text())["files"])
    ds = _create_dataset(gcp_db, access, payload.name, payload.company_id, filename, result)
    shutil.rmtree(folder, ignore_errors=True)
    return {"id": ds.id, "report": result.report()}


class SampleRequest(BaseModel):
    company_id: Optional[str] = None


@router.post("/datasets/sample", status_code=201)
def load_sample(payload: SampleRequest, access: Access = Depends(maintenance_access), gcp_db: Session = Depends(get_gcp_db)):
    """Import the 100 % fictive demonstration workbook through the regular import path."""
    sources = importer.read_sources([(SAMPLE_FILE.name, SAMPLE_FILE.read_bytes())])
    pv = importer.preview(sources)
    result = importer.build(sources, {t: {"source": v["source"], "mapping": v["mapping"]} for t, v in pv["tables"].items()})
    ds = _create_dataset(gcp_db, access, "Jeu de démonstration fictif", payload.company_id, SAMPLE_FILE.name, result)
    return {"id": ds.id, "report": result.report()}


# ── Cockpit ───────────────────────────────────────────────────────────────────
@router.get("/datasets/{dataset_id}/overview")
def overview(dataset_id: str, access: Access = Depends(maintenance_access), gcp_db: Session = Depends(get_gcp_db)):
    ds = get_dataset(dataset_id, access, gcp_db)
    eng = load_engine(ds, gcp_db)
    t = eng.study_date()
    scored = {s.notice.code: s for s in eng.score_all(t)}
    decisions = latest_decisions(ds.id, gcp_db, "notice")
    flux_counts = {f: {"notices": sum(1 for n in eng.notices if n.flux == f), "work_orders": sum(1 for o in eng.work_orders if o.flux == f)}
                   for f in ("ROUGE", "VERT", "JAUNE")}
    return {
        "study_date": t.isoformat(),
        "flux": flux_counts,
        "priorities": {p: sum(1 for s in scored.values() if s.priority == p) for p in ("HIGH", "MEDIUM", "LOW")},
        "open_high": sum(1 for s in scored.values() if s.priority == "HIGH" and not s.notice.is_closed),
        "decisions": len(decisions),
        "equipments": [equipment_json(eng, code) for code in sorted(eng.equipments)],
        "events": [
            *[{"kind": "notice", "code": n.code, "date": n.date.isoformat(), "equipment": n.equipment, "flux": n.flux,
               "symptom": n.symptom, "score": scored[n.code].score if n.code in scored else None,
               "priority": scored[n.code].priority if n.code in scored else None} for n in eng.notices],
            *[{"kind": "work_order", "code": o.code, "date": o.date.isoformat(), "equipment": o.equipment, "flux": o.flux,
               "symptom": o.symptom, "downtime_hours": o.downtime_hours} for o in eng.work_orders],
        ],
    }


@router.get("/datasets/{dataset_id}/notices")
def list_notices(dataset_id: str, as_of: Optional[str] = None, access: Access = Depends(maintenance_access),
                 gcp_db: Session = Depends(get_gcp_db)):
    ds = get_dataset(dataset_id, access, gcp_db)
    eng = load_engine(ds, gcp_db)
    t = parse_as_of(as_of, eng)
    decisions = latest_decisions(ds.id, gcp_db, "notice")
    rows = []
    for rank, s in enumerate(eng.score_all(t), start=1):
        n = s.notice
        eq = equipment_json(eng, n.equipment)
        rows.append({"rank": rank, "code": n.code, "date": n.date.isoformat(), "equipment": n.equipment, "equipment_tag": eq["tag"],
                     "equipment_name": eq["name"], "symptom": n.symptom, "category": n.category, "status": n.status,
                     "closed": n.is_closed, "postponements": n.postponements, "score": s.score, "priority": s.priority,
                     "notes": s.notes, "overridden": [c for c, ev in s.evidence.items() if ev.overridden],
                     "drivers": [c for c in E.CRITERIA if s.notes[c] >= 4],
                     "decision": decision_json(decisions.get(n.code))})
    return {"as_of": t.isoformat(), "thresholds": eng.config.thresholds, "weights": eng.config.weights, "notices": rows}


@router.get("/datasets/{dataset_id}/notices/{code}")
def notice_detail(dataset_id: str, code: str, as_of: Optional[str] = None, lang: str = "fr",
                  access: Access = Depends(maintenance_access), gcp_db: Session = Depends(get_gcp_db)):
    ds = get_dataset(dataset_id, access, gcp_db)
    eng = load_engine(ds, gcp_db)
    n = eng.notice_by_code.get(code)
    if n is None or n.flux != "JAUNE":
        raise HTTPException(status_code=404, detail="Yellow notice not found.")
    t = parse_as_of(as_of, eng)
    s = eng.score_notice(n, t)
    same_eq_reds = [o for o in eng.reds if o.equipment == n.equipment]
    links = []
    for o in eng.reds:
        for c in eng.candidate_links(o):
            if c["notice"] == n.code:
                links.append({**c, "source": "engine"})
    imported = gcp_db.query(models.MaintLink).filter_by(dataset_id=ds.id, notice_code=n.code).all()
    for l in imported:
        match = next((x for x in links if x["work_order"] == l.wo_code), None)
        if match:
            match["imported_status"] = l.expert_status
        else:
            wo = eng.wo_by_code.get(l.wo_code)
            links.append({"notice": n.code, "work_order": l.wo_code, "lead_days": (wo.date - n.date).days if wo else None,
                          "confidence": None, "same_equipment": bool(wo and wo.equipment == n.equipment),
                          "shared_families": [], "source": "import", "imported_status": l.expert_status})
    link_decisions = latest_decisions(ds.id, gcp_db, "link")
    for l in links:
        l["decision"] = decision_json(link_decisions.get(f"{l['notice']}>{l['work_order']}"))
    history = gcp_db.query(models.MaintDecision).filter(models.MaintDecision.dataset_id == ds.id,
                                                       models.MaintDecision.target.in_([n.code, *[f"{n.code}>{l['work_order']}" for l in links]])) \
        .order_by(models.MaintDecision.created_at.desc(), models.MaintDecision.id.desc()).all()
    wo_json = lambda o: {"code": o.code, "date": o.date.isoformat(), "symptom": o.symptom, "cause": o.cause, "downtime_hours": o.downtime_hours}
    return {
        "as_of": t.isoformat(),
        "notice": {"code": n.code, "date": n.date.isoformat(), "equipment": equipment_json(eng, n.equipment), "symptom": n.symptom,
                   "description": n.description, "category": n.category, "status": n.status, "postponements": n.postponements,
                   "issuer": n.issuer, "context_note": n.context_note, "overrides": n.overrides,
                   "families": sorted(eng.notice_families(n))},
        "score": s.score, "priority": s.priority, "notes": s.notes,
        "points": {c: round(s.notes[c] / 5 * eng.config.weights[c], 1) for c in E.CRITERIA},
        "weights": eng.config.weights,
        "evidence": {c: evidence_json(ev) for c, ev in s.evidence.items()},
        "prior_reds": [wo_json(o) for o in same_eq_reds if o.date < t],
        "later_reds": [wo_json(o) for o in same_eq_reds if o.date >= t],
        "links": links,
        "checks": E.suggested_checks(n, "en" if lang == "en" else "fr"),
        "history": [decision_json(d) for d in history],
    }


class NotesUpdate(BaseModel):
    overrides: dict[str, Optional[int]] = Field(default_factory=dict)
    context_note: Optional[int] = Field(default=None, ge=0, le=5)
    clear_context: bool = False
    comment: Optional[str] = None


@router.put("/datasets/{dataset_id}/notices/{code}/notes")
def update_notes(dataset_id: str, code: str, payload: NotesUpdate, access: Access = Depends(maintenance_access),
                 gcp_db: Session = Depends(get_gcp_db)):
    """Expert correction: manual context note and/or overrides of computed notes (logged)."""
    ds = get_dataset(dataset_id, access, gcp_db)
    row = gcp_db.query(models.MaintNotice).filter_by(dataset_id=ds.id, code=code, flux="JAUNE").first()
    if row is None:
        raise HTTPException(status_code=404, detail="Yellow notice not found.")
    for crit, value in payload.overrides.items():
        if crit not in E.CRITERIA or crit == "context":
            raise HTTPException(status_code=400, detail=f"Unknown criterion: {crit}")
        if value is not None and not 0 <= value <= 5:
            raise HTTPException(status_code=400, detail="Notes must be between 0 and 5.")
    before = {"context_note": row.context_note, "overrides": dict(row.overrides or {})}
    overrides = dict(row.overrides or {})
    for crit, value in payload.overrides.items():
        if value is None:
            overrides.pop(crit, None)
        else:
            overrides[crit] = value
    row.overrides = overrides or None
    if payload.clear_context:
        row.context_note = None
    elif payload.context_note is not None:
        row.context_note = payload.context_note
    gcp_db.flush()
    eng = load_engine(ds, gcp_db)
    score = eng.score_notice(eng.notice_by_code[code], eng.study_date()).score
    log_decision(gcp_db, ds, access, "notice", code, "corrected", payload.comment,
                 {"before": before, "after": {"context_note": row.context_note, "overrides": row.overrides or {}}}, score)
    gcp_db.commit()
    return {"score": score}


class DecisionCreate(BaseModel):
    target_type: Literal["notice", "link"]
    target: str = Field(min_length=1, max_length=210)
    decision: Literal["validated", "rejected"]
    comment: Optional[str] = Field(default=None, max_length=2000)


@router.post("/datasets/{dataset_id}/decisions", status_code=201)
def create_decision(dataset_id: str, payload: DecisionCreate, access: Access = Depends(maintenance_access),
                    gcp_db: Session = Depends(get_gcp_db)):
    ds = get_dataset(dataset_id, access, gcp_db)
    eng = load_engine(ds, gcp_db)
    score = None
    if payload.target_type == "notice":
        n = eng.notice_by_code.get(payload.target)
        if n is None or n.flux != "JAUNE":
            raise HTTPException(status_code=404, detail="Yellow notice not found.")
        score = eng.score_notice(n, eng.study_date()).score
    else:
        notice, _, wo = payload.target.partition(">")
        if notice not in eng.notice_by_code or wo not in eng.wo_by_code:
            raise HTTPException(status_code=404, detail="Link not found.")
    row = log_decision(gcp_db, ds, access, payload.target_type, payload.target, payload.decision, payload.comment, None, score)
    gcp_db.commit()
    return decision_json(row)


@router.get("/datasets/{dataset_id}/decisions")
def list_decisions(dataset_id: str, access: Access = Depends(maintenance_access), gcp_db: Session = Depends(get_gcp_db)):
    ds = get_dataset(dataset_id, access, gcp_db)
    rows = gcp_db.query(models.MaintDecision).filter_by(dataset_id=ds.id) \
        .order_by(models.MaintDecision.created_at.desc(), models.MaintDecision.id.desc()).all()
    return [decision_json(r) for r in rows]


# ── Replay & backtest ─────────────────────────────────────────────────────────
@router.get("/datasets/{dataset_id}/replay")
def replay_index(dataset_id: str, access: Access = Depends(maintenance_access), gcp_db: Session = Depends(get_gcp_db)):
    ds = get_dataset(dataset_id, access, gcp_db)
    eng = load_engine(ds, gcp_db)
    return [{"code": o.code, "date": o.date.isoformat(), "equipment": equipment_json(eng, o.equipment)["tag"], "symptom": o.symptom,
             "candidates": len(eng.candidate_links(o))} for o in eng.reds]


@router.get("/datasets/{dataset_id}/replay/{wo_code}")
def replay(dataset_id: str, wo_code: str, threshold: Optional[int] = None, access: Access = Depends(maintenance_access),
           gcp_db: Session = Depends(get_gcp_db)):
    ds = get_dataset(dataset_id, access, gcp_db)
    eng = load_engine(ds, gcp_db)
    wo = eng.wo_by_code.get(wo_code)
    if wo is None or wo.flux != "ROUGE":
        raise HTTPException(status_code=404, detail="Red work order not found.")
    result = eng.replay(wo, threshold)
    link_decisions = latest_decisions(ds.id, gcp_db, "link")
    imported = {l.notice_code: l.expert_status for l in gcp_db.query(models.MaintLink).filter_by(dataset_id=ds.id, wo_code=wo_code).all()}
    for p in result["precursors"]:
        n = eng.notice_by_code[p["notice"]]
        p.update({"date": n.date.isoformat(), "symptom": n.symptom, "equipment_tag": equipment_json(eng, n.equipment)["tag"],
                  "imported_status": imported.get(n.code),
                  "decision": decision_json(link_decisions.get(f"{n.code}>{wo_code}"))})
    result["work_order"] = {"code": wo.code, "date": wo.date.isoformat(), "type": wo.type, "symptom": wo.symptom, "cause": wo.cause,
                            "downtime_hours": wo.downtime_hours, "equipment": equipment_json(eng, wo.equipment)}
    return result


@router.get("/datasets/{dataset_id}/backtest")
def backtest(dataset_id: str, threshold: Optional[int] = None, horizon: Optional[int] = None,
             access: Access = Depends(maintenance_access), gcp_db: Session = Depends(get_gcp_db)):
    ds = get_dataset(dataset_id, access, gcp_db)
    eng = load_engine(ds, gcp_db)
    if threshold is not None and not 1 <= threshold <= 100:
        raise HTTPException(status_code=400, detail="threshold must be between 1 and 100")
    if horizon is not None and not 1 <= horizon <= 730:
        raise HTTPException(status_code=400, detail="horizon must be between 1 and 730 days")
    return eng.backtest(threshold, horizon)


# ── Export ────────────────────────────────────────────────────────────────────
@router.get("/datasets/{dataset_id}/export.xlsx")
def export_xlsx(dataset_id: str, access: Access = Depends(maintenance_access), gcp_db: Session = Depends(get_gcp_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    ds = get_dataset(dataset_id, access, gcp_db)
    eng = load_engine(ds, gcp_db)
    t = eng.study_date()
    decisions = latest_decisions(ds.id, gcp_db, "notice")
    wb = Workbook()

    ws = wb.active
    ws.title = "Priorités"
    ws.append(["Rang", "Avis", "Date", "Équipement", "Symptôme", "Statut", "Reports", *E.CRITERIA, "Score /100", "Classe",
               "Décision", "Commentaire", "Par"])
    for rank, s in enumerate(eng.score_all(t), start=1):
        n, d = s.notice, decisions.get(s.notice.code)
        ws.append([rank, n.code, n.date, equipment_json(eng, n.equipment)["tag"], n.symptom, n.status, n.postponements,
                   *[s.notes[c] for c in E.CRITERIA], s.score, s.priority,
                   d.decision if d else "", (d.comment or "") if d else "", d.actor_name if d else ""])

    ws2 = wb.create_sheet("Décisions")
    ws2.append(["Date", "Type", "Cible", "Décision", "Commentaire", "Score", "Par", "Détail"])
    for d in gcp_db.query(models.MaintDecision).filter_by(dataset_id=ds.id).order_by(models.MaintDecision.created_at).all():
        ws2.append([d.created_at, d.target_type, d.target, d.decision, d.comment or "", d.score_snapshot, d.actor_name,
                    json.dumps(d.changes, ensure_ascii=False) if d.changes else ""])

    ws3 = wb.create_sheet("Paramètres")
    ws3.append(["Jeu de données", ds.name])
    ws3.append(["Date d’étude", t])
    for c in E.CRITERIA:
        ws3.append([f"Poids — {c}", ds.weights.get(c)])
    for k, v in (ds.thresholds or {}).items():
        ws3.append([f"Seuil — {k}", v])
    for k, v in (ds.windows or {}).items():
        ws3.append([f"Fenêtre — {k}", v])
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = "".join(ch for ch in ds.name if ch.isalnum() or ch in " -_").strip().replace(" ", "_") or "strap_maintenance"
    ascii_name = name.encode("ascii", "ignore").decode() or "strap_maintenance"
    disposition = f'attachment; filename="{ascii_name}.xlsx"; filename*=UTF-8\'\'{quote(name)}.xlsx'
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": disposition})
